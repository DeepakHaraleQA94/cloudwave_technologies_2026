from dotenv import load_dotenv
from pathlib import Path
import os

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

from fastapi import FastAPI, APIRouter, HTTPException, Depends, Request, UploadFile, File, Query
from fastapi.responses import StreamingResponse, Response
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
from pydantic import BaseModel, EmailStr, Field
from typing import List, Optional, Any, Dict
from datetime import datetime, timezone, timedelta
import logging, uuid, re, io, csv, base64, secrets, asyncio, ipaddress
import bcrypt, jwt, httpx
from html import escape
from html.parser import HTMLParser
from urllib.parse import urlparse

# ------------------------------------------------------------------ DB
mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

JWT_SECRET = os.environ['JWT_SECRET']
JWT_ALG = "HS256"

app = FastAPI(title="CloudWave Technologies API")
api = APIRouter(prefix="/api")

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger("cloudwave")


# ------------------------------------------------------------------ helpers
def now_iso():
    return datetime.now(timezone.utc).isoformat()

def new_id():
    return str(uuid.uuid4())

def slugify(text: str) -> str:
    s = re.sub(r'[^a-z0-9]+', '-', (text or '').lower()).strip('-')
    return s or new_id()[:8]

def hash_password(p: str) -> str:
    return bcrypt.hashpw(p.encode(), bcrypt.gensalt()).decode()

def verify_password(p: str, h: str) -> bool:
    try:
        return bcrypt.checkpw(p.encode(), h.encode())
    except Exception:
        return False

def create_token(user_id: str, email: str) -> str:
    payload = {"sub": user_id, "email": email,
               "exp": datetime.now(timezone.utc) + timedelta(days=7), "type": "access"}
    return jwt.encode(payload, JWT_SECRET, algorithm=JWT_ALG)

async def get_current_admin(request: Request) -> dict:
    auth = request.headers.get("Authorization", "")
    token = auth[7:] if auth.startswith("Bearer ") else request.cookies.get("access_token")
    if not token:
        raise HTTPException(401, "Not authenticated")
    try:
        payload = jwt.decode(token, JWT_SECRET, algorithms=[JWT_ALG])
    except jwt.ExpiredSignatureError:
        raise HTTPException(401, "Token expired")
    except jwt.InvalidTokenError:
        raise HTTPException(401, "Invalid token")
    user = await db.users.find_one({"id": payload["sub"]}, {"_id": 0, "password_hash": 0})
    if not user:
        raise HTTPException(401, "User not found")
    return user


# ------------------------------------------------------------------ auth models
class LoginIn(BaseModel):
    email: EmailStr
    password: str

class ForgotIn(BaseModel):
    email: EmailStr

class ResetIn(BaseModel):
    token: str
    password: str


@api.post("/auth/login")
async def login(body: LoginIn):
    email = body.email.lower()
    user = await db.users.find_one({"email": email})
    if not user or not verify_password(body.password, user["password_hash"]):
        raise HTTPException(401, "Invalid email or password")
    token = create_token(user["id"], email)
    return {"token": token, "user": {"id": user["id"], "email": email,
            "name": user.get("name"), "role": user.get("role", "admin")}}

@api.get("/auth/me")
async def me(admin=Depends(get_current_admin)):
    return admin

@api.post("/auth/logout")
async def logout(admin=Depends(get_current_admin)):
    return {"ok": True}

@api.post("/auth/forgot-password")
async def forgot(body: ForgotIn):
    user = await db.users.find_one({"email": body.email.lower()})
    if user:
        tok = secrets.token_urlsafe(32)
        await db.password_reset_tokens.insert_one({
            "token": tok, "user_id": user["id"], "used": False,
            "expires_at": (datetime.now(timezone.utc) + timedelta(hours=1)).isoformat()})
        logger.info(f"[PASSWORD RESET] link token for {body.email}: {tok}")
    return {"ok": True, "message": "If the email exists, a reset link has been sent."}

@api.post("/auth/reset-password")
async def reset(body: ResetIn):
    rec = await db.password_reset_tokens.find_one({"token": body.token, "used": False})
    if not rec or rec["expires_at"] < now_iso():
        raise HTTPException(400, "Invalid or expired token")
    await db.users.update_one({"id": rec["user_id"]},
                              {"$set": {"password_hash": hash_password(body.password)}})
    await db.password_reset_tokens.update_one({"token": body.token}, {"$set": {"used": True}})
    return {"ok": True}


# ------------------------------------------------------------------ media upload (base64 stored in mongo)
ALLOWED_IMG = {"image/jpeg", "image/jpg", "image/png", "image/webp", "image/gif"}
BACKEND_PUBLIC = os.environ.get("BACKEND_PUBLIC_URL", "")

@api.post("/admin/upload")
async def upload_image(file: UploadFile = File(...), admin=Depends(get_current_admin)):
    if file.content_type not in ALLOWED_IMG:
        raise HTTPException(400, "Only JPG, PNG, WEBP, GIF images are allowed")
    data = await file.read()
    if len(data) > 6 * 1024 * 1024:
        raise HTTPException(400, "File too large (max 6MB)")
    mid = new_id()
    await db.media_files.insert_one({
        "id": mid, "content_type": file.content_type,
        "data": base64.b64encode(data).decode(), "created_at": now_iso()})
    return {"url": f"/api/media/{mid}", "id": mid}

@api.get("/media/{mid}")
async def get_media(mid: str):
    rec = await db.media_files.find_one({"id": mid}, {"_id": 0})
    if not rec:
        raise HTTPException(404, "Not found")
    if rec.get("private"):
        raise HTTPException(403, "This file is private")
    return Response(content=base64.b64decode(rec["data"]),
                    media_type=rec["content_type"],
                    headers={"Cache-Control": "public, max-age=31536000"})


# ------------------------------------------------------------------ EMAIL (Emergent-managed Resend)
EMAIL_BASE_URL = "https://integrations.emergentagent.com"
EMAIL_KEY = os.environ.get("EMERGENT_EMAIL_KEY")
EMAIL_FROM_NAME = os.environ.get("EMAIL_FROM_NAME", "CloudWave Technologies")

class _EmailScan(HTMLParser):
    def __init__(self):
        super().__init__()
        self.tags, self.urls = set(), []
    def handle_starttag(self, tag, attrs):
        self.tags.add(tag.lower())
        self.urls += [v for k, v in attrs if k.lower() in ("href", "src") and v]

def _assert_safe_email(subject: str, html: str) -> None:
    scan = _EmailScan(); scan.feed(html)
    if scan.tags & {"form", "input", "textarea", "select"}:
        raise ValueError("No forms or input fields in email")
    for url in scan.urls:
        low = url.strip().lower()
        if low.startswith(("mailto:", "tel:", "cid:", "#")):
            continue
        if not low.startswith("https://"):
            raise ValueError("Email links/assets must be absolute https")

async def send_email(*, to: str, subject: str, html: str):
    _assert_safe_email(subject, html)
    payload = {"to": [to], "subject": subject, "html": html, "from_name": EMAIL_FROM_NAME}
    async with httpx.AsyncClient(timeout=30) as c:
        resp = await c.post(f"{EMAIL_BASE_URL}/api/v1/email/send",
                            headers={"X-Email-Key": EMAIL_KEY}, json=payload)
    resp.raise_for_status()
    return resp.json().get("id")

async def notify_admins_new_enquiry(enq: dict):
    if not EMAIL_KEY:
        return
    recips = [e.strip() for e in os.environ.get("NOTIFY_EMAILS", "").split(",") if e.strip()]
    if not recips:
        admins = await db.users.find({"role": "admin"}, {"_id": 0, "email": 1}).to_list(50)
        recips = [a["email"] for a in admins]
    subject = f"New Enquiry: {enq.get('name')} — {enq.get('course_name') or 'General'}"
    rows = [("Enquiry ID", enq.get("enquiry_id")), ("Name", enq.get("name")),
            ("Email", enq.get("email")), ("Mobile", enq.get("mobile")),
            ("WhatsApp", enq.get("whatsapp")), ("Course", enq.get("course_name")),
            ("Batch", enq.get("batch_name")), ("Preferred Mode", enq.get("mode")),
            ("City", enq.get("city")), ("Source", enq.get("source")),
            ("Message", enq.get("message"))]
    tr = "".join(f'<tr><td style="padding:6px 12px;color:#64748b;font-weight:600">{escape(str(k))}</td>'
                 f'<td style="padding:6px 12px;color:#0f172a">{escape(str(v or "-"))}</td></tr>' for k, v in rows)
    html = (f'<table role="presentation" width="100%" style="font-family:Arial,sans-serif">'
            f'<tr><td style="padding:20px">'
            f'<h2 style="color:#1D4ED8;margin:0 0 4px">New Student Enquiry</h2>'
            f'<p style="color:#475569;margin:0 0 16px">A new enquiry was submitted on the CloudWave Technologies website.</p>'
            f'<table role="presentation" width="100%" style="border:1px solid #e2e8f0;border-radius:8px;border-collapse:separate">{tr}</table>'
            f'<p style="font-size:12px;color:#94a3b8;margin-top:16px">Sent by {escape(EMAIL_FROM_NAME)}. '
            f'Sign in to your admin dashboard to view and manage this enquiry. '
            f'We never ask for your password or payment details by email.</p>'
            f'</td></tr></table>')
    for r in recips:
        try:
            await send_email(to=r, subject=subject, html=html)
        except Exception as e:
            logger.error(f"Enquiry notification email failed for {r}: {e}")


async def send_student_confirmation(enq: dict):
    if not EMAIL_KEY or not enq.get("email"):
        return
    subject = f"We received your enquiry — {EMAIL_FROM_NAME}"
    course = enq.get("course_name") or "our courses"
    html = (f'<table role="presentation" width="100%" style="font-family:Arial,sans-serif">'
            f'<tr><td style="padding:24px">'
            f'<h2 style="color:#1D4ED8;margin:0 0 6px">Thank you, {escape(str(enq.get("name") or "there"))}!</h2>'
            f'<p style="color:#475569;margin:0 0 16px;line-height:1.6">We have received your enquiry about '
            f'<strong>{escape(str(course))}</strong>. Our counsellors will reach out to you very shortly with '
            f'course details, fees and upcoming batch schedules.</p>'
            f'<div style="background:#f1f5f9;border-radius:8px;padding:14px 16px;margin:0 0 16px">'
            f'<p style="margin:0;color:#0f172a">Your Enquiry ID: <strong style="color:#1D4ED8">{escape(str(enq.get("enquiry_id")))}</strong></p>'
            f'<p style="margin:6px 0 0;color:#64748b;font-size:13px">Please quote this ID in any follow-up conversation.</p>'
            f'</div>'
            f'<p style="color:#475569;margin:0 0 4px;line-height:1.6">If you need immediate help, just reply to this email '
            f'or contact us on WhatsApp.</p>'
            f'<p style="font-size:12px;color:#94a3b8;margin-top:16px">Sent by {escape(EMAIL_FROM_NAME)}. '
            f'This is a confirmation of your enquiry — we never ask for your password or payment details by email.</p>'
            f'</td></tr></table>')
    try:
        await send_email(to=enq["email"], subject=subject, html=html)
    except Exception as e:
        logger.error(f"Student confirmation email failed for {enq.get('email')}: {e}")


# ------------------------------------------------------------------ WhatsApp (Twilio) alerts
TWILIO_SID = os.environ.get("TWILIO_ACCOUNT_SID", "").strip()
TWILIO_TOKEN = os.environ.get("TWILIO_AUTH_TOKEN", "").strip()
TWILIO_WA_FROM = os.environ.get("TWILIO_WHATSAPP_FROM", "").strip()
ADMIN_WA_NUMBERS = [n.strip() for n in os.environ.get("ADMIN_WHATSAPP_NUMBERS", "").split(",") if n.strip()]

def is_high_intent(doc: dict) -> bool:
    has_course = bool(doc.get("course_id") or doc.get("course_name"))
    has_batch = bool(doc.get("batch_id"))
    msg = doc.get("message") or ""
    return bool(has_course and (has_batch or len(msg.strip()) >= 15))

async def send_whatsapp(to: str, body: str):
    url = f"https://api.twilio.com/2010-04-01/Accounts/{TWILIO_SID}/Messages.json"
    data = {"From": TWILIO_WA_FROM, "To": f"whatsapp:{to}", "Body": body}
    async with httpx.AsyncClient(timeout=15) as c:
        r = await c.post(url, data=data, auth=(TWILIO_SID, TWILIO_TOKEN))
    r.raise_for_status()
    return r.json().get("sid")

async def notify_admins_whatsapp(enq: dict):
    if not (TWILIO_SID and TWILIO_TOKEN and TWILIO_WA_FROM and ADMIN_WA_NUMBERS):
        logger.info("WhatsApp alerts not configured (Twilio env vars missing) — skipping.")
        return
    body = ("*High-intent enquiry* \U0001F525\n"
            f"ID: {enq.get('enquiry_id')}\n"
            f"Name: {enq.get('name')}\n"
            f"Mobile: {enq.get('mobile')}\n"
            f"Course: {enq.get('course_name') or '-'}\n"
            f"Batch: {enq.get('batch_name') or '-'}\n"
            f"City: {enq.get('city') or '-'}\n"
            "— CloudWave Technologies")
    for n in ADMIN_WA_NUMBERS:
        try:
            await send_whatsapp(n, body)
        except Exception as e:
            logger.error(f"WhatsApp alert failed for {n}: {e}")


# ------------------------------------------------------------------ generic CRUD factory
# resource -> (collection, published_field)
RESOURCES = {
    "courses": ("courses", "published"),
    "batches": ("batches", "published"),
    "trainers": ("trainers", "active"),
    "testimonials": ("testimonials", "published"),
    "faqs": ("faqs", "published"),
    "blog": ("blog_posts", "published"),
    "gallery": ("gallery", "published"),
    "videos": ("videos", "published"),
    "events": ("events", "published"),
    "placements": ("placements", "published"),
    "themes": ("themes", None),
    "certificates": ("certificates", "published"),
    "expenses": ("expenses", None),
    "slides": ("slides", "published"),
    "technologies": ("technologies", "active"),
    "learning_resources": ("learning_resources", None),
}

def clean(doc):
    doc.pop("_id", None)
    return doc

@api.get("/admin/data/{resource}")
async def admin_list(resource: str, admin=Depends(get_current_admin)):
    if resource not in RESOURCES:
        raise HTTPException(404, "Unknown resource")
    col = RESOURCES[resource][0]
    items = await db[col].find({}, {"_id": 0}).sort("sort_order", 1).to_list(2000)
    items.sort(key=lambda x: (x.get("sort_order", 0), x.get("created_at", "")))
    return items

@api.post("/admin/data/{resource}")
async def admin_create(resource: str, body: Dict[str, Any], admin=Depends(get_current_admin)):
    if resource not in RESOURCES:
        raise HTTPException(404, "Unknown resource")
    col = RESOURCES[resource][0]
    body["id"] = new_id()
    body["created_at"] = now_iso()
    body["updated_at"] = now_iso()
    if "sort_order" not in body:
        body["sort_order"] = await db[col].count_documents({})
    if resource in ("courses", "blog") and body.get("name" if resource == "courses" else "title"):
        base = body.get("name") if resource == "courses" else body.get("title")
        body["slug"] = body.get("slug") or slugify(base)
    await db[col].insert_one(dict(body))
    return clean(body)

@api.put("/admin/data/{resource}/{item_id}")
async def admin_update(resource: str, item_id: str, body: Dict[str, Any], admin=Depends(get_current_admin)):
    if resource not in RESOURCES:
        raise HTTPException(404, "Unknown resource")
    col = RESOURCES[resource][0]
    body.pop("id", None); body.pop("_id", None)
    body["updated_at"] = now_iso()
    res = await db[col].update_one({"id": item_id}, {"$set": body})
    if res.matched_count == 0:
        raise HTTPException(404, "Not found")
    doc = await db[col].find_one({"id": item_id}, {"_id": 0})
    return doc

@api.delete("/admin/data/{resource}/{item_id}")
async def admin_delete(resource: str, item_id: str, admin=Depends(get_current_admin)):
    if resource not in RESOURCES:
        raise HTTPException(404, "Unknown resource")
    col = RESOURCES[resource][0]
    await db[col].delete_one({"id": item_id})
    return {"ok": True}


# ------------------------------------------------------------------ PUBLIC read endpoints
async def pub_list(col, pub_field, extra=None):
    q = {}
    if pub_field:
        q[pub_field] = True
    if extra:
        q.update(extra)
    items = await db[col].find(q, {"_id": 0}).to_list(2000)
    items.sort(key=lambda x: (x.get("sort_order", 0), x.get("created_at", "")))
    return items

@api.get("/settings")
async def get_settings():
    s = await db.website_settings.find_one({"id": "main"}, {"_id": 0})
    return s or {}

@api.put("/admin/settings")
async def update_settings(body: Dict[str, Any], admin=Depends(get_current_admin)):
    body["id"] = "main"; body["updated_at"] = now_iso()
    await db.website_settings.update_one({"id": "main"}, {"$set": body}, upsert=True)
    return await db.website_settings.find_one({"id": "main"}, {"_id": 0})

@api.get("/courses")
async def public_courses(category: Optional[str] = None, featured: Optional[bool] = None):
    extra = {}
    if category: extra["category"] = category
    if featured is not None: extra["featured"] = featured
    return await pub_list("courses", "published", extra)

@api.get("/courses/{slug}")
async def public_course(slug: str):
    c = await db.courses.find_one({"slug": slug, "published": True}, {"_id": 0})
    if not c:
        c = await db.courses.find_one({"id": slug}, {"_id": 0})
    if not c:
        raise HTTPException(404, "Course not found")
    return c

@api.get("/batches")
async def public_batches(course_id: Optional[str] = None):
    extra = {"course_id": course_id} if course_id else None
    return await pub_list("batches", "published", extra)

@api.get("/trainers")
async def public_trainers():
    return await pub_list("trainers", "active")

@api.get("/testimonials")
async def public_testimonials():
    return await pub_list("testimonials", "published")

@api.get("/faqs")
async def public_faqs(category: Optional[str] = None):
    extra = {"category": category} if category else None
    return await pub_list("faqs", "published", extra)

@api.get("/blog")
async def public_blog(category: Optional[str] = None):
    extra = {"category": category} if category else None
    return await pub_list("blog_posts", "published", extra)

@api.get("/blog/{slug}")
async def public_blog_post(slug: str):
    p = await db.blog_posts.find_one({"slug": slug, "published": True}, {"_id": 0})
    if not p:
        raise HTTPException(404, "Post not found")
    return p

@api.get("/gallery")
async def public_gallery(category: Optional[str] = None, year: Optional[str] = None):
    extra = {}
    if category and category != "All": extra["category"] = category
    if year: extra["year"] = year
    return await pub_list("gallery", "published", extra)

@api.get("/videos")
async def public_videos(category: Optional[str] = None):
    extra = {"category": category} if category else None
    return await pub_list("videos", "published", extra)

@api.get("/events")
async def public_events():
    return await pub_list("events", "published")

@api.get("/events/{item_id}")
async def public_event(item_id: str):
    e = await db.events.find_one({"id": item_id, "published": True}, {"_id": 0})
    if not e:
        raise HTTPException(404, "Event not found")
    return e

@api.get("/placements")
async def public_placements(course: Optional[str] = None, company: Optional[str] = None, year: Optional[str] = None):
    extra = {}
    if course: extra["course"] = course
    if company: extra["company_name"] = company
    if year: extra["placement_year"] = year
    return await pub_list("placements", "published", extra)

DAY_NAMES = ["monday", "tuesday", "wednesday", "thursday", "friday", "saturday", "sunday"]
TECH_PRESENTATION_DEFAULTS = {
    "tech_style": "Float", "tech_speed": 5, "tech_direction": "left",
    "tech_delay": 150, "tech_loop": True, "tech_hover": "3d-tilt",
}

@api.get("/technologies")
async def public_technologies():
    today = DAY_NAMES[datetime.now().weekday()]
    techs = await db.technologies.find({"active": True}, {"_id": 0}).to_list(500)
    techs.sort(key=lambda x: (x.get("sort_order", 0), x.get("created_at", "")))
    shown = [t for t in techs if (t.get("day_theme") or "everyday") in ("everyday", today)]
    if not shown:
        shown = techs
    s = await db.website_settings.find_one({"id": "main"}, {"_id": 0}) or {}
    pres = {k: (s.get(k) if s.get(k) is not None else v) for k, v in TECH_PRESENTATION_DEFAULTS.items()}
    return {"technologies": shown, "presentation": pres}

@api.get("/certificates/verify")
async def verify_certificate(cid: str):
    c = await db.certificates.find_one(
        {"certificate_id": {"$regex": f"^{re.escape(cid.strip())}$", "$options": "i"}, "published": True},
        {"_id": 0})
    if not c:
        raise HTTPException(404, "No valid certificate found for this ID")
    return c

@api.get("/certificates/{cid}/download")
async def download_certificate(cid: str):
    c = await db.certificates.find_one(
        {"certificate_id": {"$regex": f"^{re.escape(cid.strip())}$", "$options": "i"}, "published": True},
        {"_id": 0})
    if not c:
        raise HTTPException(404, "No valid certificate found for this ID")
    from reportlab.lib.pagesizes import landscape, A4
    from reportlab.pdfgen import canvas
    from reportlab.lib import colors

    settings = await db.website_settings.find_one({"id": "main"}, {"_id": 0}) or {}
    inst = settings.get("institute_name", "CloudWave Technologies")
    blue = colors.HexColor("#1D4ED8")
    accent = colors.HexColor("#F97316")
    dark = colors.HexColor("#0F172A")
    muted = colors.HexColor("#64748B")

    buf = io.BytesIO()
    W, H = landscape(A4)
    pdf = canvas.Canvas(buf, pagesize=landscape(A4))

    pdf.setFillColor(colors.HexColor("#F8FAFC"))
    pdf.rect(0, 0, W, H, fill=1, stroke=0)
    pdf.setStrokeColor(blue); pdf.setLineWidth(6)
    pdf.rect(24, 24, W - 48, H - 48, fill=0, stroke=1)
    pdf.setStrokeColor(accent); pdf.setLineWidth(1.5)
    pdf.rect(36, 36, W - 72, H - 72, fill=0, stroke=1)

    cx = W / 2
    pdf.setFillColor(blue)
    pdf.setFont("Helvetica-Bold", 26)
    pdf.drawCentredString(cx, H - 100, inst)
    pdf.setFillColor(muted)
    pdf.setFont("Helvetica", 11)
    pdf.drawCentredString(cx, H - 120, (settings.get("tagline") or "IT Training Institute"))

    pdf.setFillColor(dark)
    pdf.setFont("Helvetica-Bold", 34)
    pdf.drawCentredString(cx, H - 185, "Certificate of Completion")
    pdf.setStrokeColor(accent); pdf.setLineWidth(2)
    pdf.line(cx - 90, H - 198, cx + 90, H - 198)

    pdf.setFillColor(muted); pdf.setFont("Helvetica", 14)
    pdf.drawCentredString(cx, H - 235, "This is to certify that")
    pdf.setFillColor(blue); pdf.setFont("Helvetica-Bold", 30)
    pdf.drawCentredString(cx, H - 275, c.get("student_name", ""))
    pdf.setFillColor(muted); pdf.setFont("Helvetica", 14)
    pdf.drawCentredString(cx, H - 305, "has successfully completed the course")
    pdf.setFillColor(dark); pdf.setFont("Helvetica-Bold", 20)
    pdf.drawCentredString(cx, H - 335, c.get("course", ""))

    details = []
    if c.get("grade"): details.append(f"Grade: {c['grade']}")
    if c.get("issue_date"): details.append(f"Issued: {c['issue_date']}")
    if details:
        pdf.setFillColor(muted); pdf.setFont("Helvetica", 12)
        pdf.drawCentredString(cx, H - 365, "     |     ".join(details))

    pdf.setFillColor(dark); pdf.setFont("Helvetica-Bold", 11)
    pdf.drawString(70, 70, f"Certificate ID: {c.get('certificate_id')}")
    pdf.setFillColor(muted); pdf.setFont("Helvetica", 9)
    pdf.drawString(70, 55, "Verify authenticity at the institute website /verify page.")
    pdf.setStrokeColor(dark); pdf.setLineWidth(1)
    pdf.line(W - 230, 78, W - 70, 78)
    pdf.setFillColor(dark); pdf.setFont("Helvetica", 10)
    pdf.drawCentredString(W - 150, 62, "Authorised Signatory")

    pdf.showPage(); pdf.save()
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{c.get("certificate_id")}.pdf"'})

@api.get("/theme/active")
async def active_theme():
    settings = await db.website_settings.find_one({"id": "main"}, {"_id": 0}) or {}
    if settings.get("theme_mode") == "off":
        return {"theme": None}
    today = datetime.now(timezone.utc).date().isoformat()
    themes = await db.themes.find({"is_active": True}, {"_id": 0}).to_list(200)
    active = []
    for t in themes:
        sd, ed = t.get("start_date"), t.get("end_date")
        if sd and sd > today:
            continue
        if ed and ed < today:
            continue
        active.append(t)
    if not active:
        return {"theme": None}
    active.sort(key=lambda x: x.get("priority", 0), reverse=True)
    return {"theme": active[0]}


# ------------------------------------------------------------------ ENQUIRIES (public submit + admin manage)
class EnquiryIn(BaseModel):
    name: str
    email: EmailStr
    mobile: str
    whatsapp: Optional[str] = ""
    course_id: Optional[str] = ""
    course_name: Optional[str] = ""
    batch_id: Optional[str] = ""
    batch_name: Optional[str] = ""
    mode: Optional[str] = ""
    city: Optional[str] = ""
    message: Optional[str] = ""
    source: Optional[str] = "Website"
    consent: bool

@api.post("/enquiries")
async def create_enquiry(body: EnquiryIn):
    if not re.match(r'^[6-9]\d{9}$', re.sub(r'\D', '', body.mobile)[-10:]):
        raise HTTPException(422, "Please enter a valid 10-digit Indian mobile number")
    if not body.consent:
        raise HTTPException(422, "Consent is required")
    count = await db.enquiries.count_documents({}) + 1
    doc = body.model_dump()
    doc["id"] = new_id()
    doc["enquiry_id"] = f"CW{datetime.now().strftime('%y%m')}{count:04d}"
    doc["status"] = "New"
    doc["follow_up_date"] = ""
    doc["notes"] = []
    doc["created_at"] = now_iso()
    doc["updated_at"] = now_iso()
    doc["high_intent"] = is_high_intent(doc)
    await db.enquiries.insert_one(dict(doc))
    asyncio.create_task(notify_admins_new_enquiry(dict(doc)))
    asyncio.create_task(send_student_confirmation(dict(doc)))
    if doc["high_intent"]:
        asyncio.create_task(notify_admins_whatsapp(dict(doc)))
    return {"ok": True, "enquiry_id": doc["enquiry_id"], "message": "Enquiry submitted successfully"}

@api.get("/admin/enquiries")
async def list_enquiries(admin=Depends(get_current_admin),
                         search: Optional[str] = None, status: Optional[str] = None,
                         course_id: Optional[str] = None, source: Optional[str] = None,
                         date_from: Optional[str] = None, date_to: Optional[str] = None,
                         page: int = 1, page_size: int = 20):
    q = {}
    if status and status != "All": q["status"] = status
    if course_id and course_id != "All": q["course_id"] = course_id
    if source and source != "All": q["source"] = source
    if date_from: q.setdefault("created_at", {})["$gte"] = date_from
    if date_to: q.setdefault("created_at", {})["$lte"] = date_to + "T23:59:59"
    if search:
        rx = {"$regex": re.escape(search), "$options": "i"}
        q["$or"] = [{"name": rx}, {"email": rx}, {"mobile": rx},
                    {"enquiry_id": rx}, {"course_name": rx}]
    total = await db.enquiries.count_documents(q)
    items = await db.enquiries.find(q, {"_id": 0}).sort("created_at", -1) \
        .skip((page - 1) * page_size).limit(page_size).to_list(page_size)
    return {"items": items, "total": total, "page": page, "page_size": page_size}

@api.put("/admin/enquiries/{item_id}")
async def update_enquiry(item_id: str, body: Dict[str, Any], admin=Depends(get_current_admin)):
    body.pop("id", None); body.pop("_id", None)
    body["updated_at"] = now_iso()
    r = await db.enquiries.update_one({"id": item_id}, {"$set": body})
    if r.matched_count == 0:
        raise HTTPException(404, "Not found")
    return await db.enquiries.find_one({"id": item_id}, {"_id": 0})

@api.post("/admin/enquiries/{item_id}/notes")
async def add_enquiry_note(item_id: str, body: Dict[str, Any], admin=Depends(get_current_admin)):
    note = {"id": new_id(), "text": body.get("text", ""),
            "author": admin.get("name", "Admin"), "created_at": now_iso()}
    r = await db.enquiries.update_one({"id": item_id}, {"$push": {"notes": note}})
    if r.matched_count == 0:
        raise HTTPException(404, "Not found")
    return note

@api.delete("/admin/enquiries/{item_id}")
async def delete_enquiry(item_id: str, admin=Depends(get_current_admin)):
    await db.enquiries.delete_one({"id": item_id})
    return {"ok": True}

@api.get("/admin/enquiries-export")
async def export_enquiries(admin=Depends(get_current_admin)):
    items = await db.enquiries.find({}, {"_id": 0}).sort("created_at", -1).to_list(10000)
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["Enquiry ID", "Name", "Email", "Mobile", "WhatsApp", "Course",
                "Batch", "Status", "Notes", "Created Date"])
    for e in items:
        notes = " | ".join(n.get("text", "") for n in e.get("notes", []))
        w.writerow([e.get("enquiry_id"), e.get("name"), e.get("email"), e.get("mobile"),
                    e.get("whatsapp"), e.get("course_name"), e.get("batch_name"),
                    e.get("status"), notes, e.get("created_at", "")[:10]])
    buf.seek(0)
    return StreamingResponse(iter([buf.getvalue()]), media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=enquiries.csv"})


# ------------------------------------------------------------------ CONTACT messages
class ContactIn(BaseModel):
    name: str
    email: EmailStr
    phone: Optional[str] = ""
    subject: Optional[str] = ""
    message: str

@api.post("/contact")
async def create_contact(body: ContactIn):
    doc = body.model_dump()
    doc["id"] = new_id(); doc["status"] = "New"; doc["created_at"] = now_iso()
    await db.contact_messages.insert_one(dict(doc))
    return {"ok": True, "message": "Message sent successfully"}

@api.get("/admin/contact-messages")
async def list_contact(admin=Depends(get_current_admin)):
    return await db.contact_messages.find({}, {"_id": 0}).sort("created_at", -1).to_list(2000)

@api.delete("/admin/contact-messages/{item_id}")
async def delete_contact(item_id: str, admin=Depends(get_current_admin)):
    await db.contact_messages.delete_one({"id": item_id})
    return {"ok": True}


# ------------------------------------------------------------------ ADMIN users management
@api.get("/admin/users")
async def list_users(admin=Depends(get_current_admin)):
    return await db.users.find({}, {"_id": 0, "password_hash": 0}).to_list(200)

@api.post("/admin/users")
async def create_user(body: Dict[str, Any], admin=Depends(get_current_admin)):
    email = body["email"].lower()
    if await db.users.find_one({"email": email}):
        raise HTTPException(400, "Email already exists")
    doc = {"id": new_id(), "email": email, "name": body.get("name", "Admin"),
           "role": body.get("role", "admin"), "password_hash": hash_password(body["password"]),
           "created_at": now_iso()}
    await db.users.insert_one(doc)
    return {"id": doc["id"], "email": email, "name": doc["name"], "role": doc["role"]}

@api.delete("/admin/users/{item_id}")
async def delete_user(item_id: str, admin=Depends(get_current_admin)):
    count = await db.users.count_documents({})
    if count <= 1:
        raise HTTPException(400, "Cannot delete the last admin")
    await db.users.delete_one({"id": item_id})
    return {"ok": True}


# ------------------------------------------------------------------ DASHBOARD stats
@api.get("/admin/stats")
async def stats(admin=Depends(get_current_admin)):
    enq = await db.enquiries.find({}, {"_id": 0}).to_list(10000)
    by_status = {}
    for e in enq:
        by_status[e.get("status", "New")] = by_status.get(e.get("status", "New"), 0) + 1
    by_course = {}
    for e in enq:
        cn = e.get("course_name") or "Not specified"
        by_course[cn] = by_course.get(cn, 0) + 1
    by_month = {}
    for e in enq:
        m = (e.get("created_at") or "")[:7]
        if m: by_month[m] = by_month.get(m, 0) + 1
    return {
        "total_enquiries": len(enq),
        "new_enquiries": by_status.get("New", 0),
        "contacted": by_status.get("Contacted", 0),
        "converted": by_status.get("Converted", 0),
        "total_courses": await db.courses.count_documents({}),
        "active_batches": await db.batches.count_documents({"status": {"$in": ["Upcoming", "Filling Fast", "Started"]}}),
        "total_trainers": await db.trainers.count_documents({}),
        "total_placements": await db.placements.count_documents({}),
        "by_status": [{"name": k, "value": v} for k, v in by_status.items()],
        "by_course": sorted([{"name": k, "value": v} for k, v in by_course.items()], key=lambda x: -x["value"])[:6],
        "by_month": [{"name": k, "value": v} for k, v in sorted(by_month.items())][-6:],
        "recent": sorted(enq, key=lambda x: x.get("created_at", ""), reverse=True)[:5],
        "total_students": await db.students.count_documents({}),
        "active_students": await db.students.count_documents({"status": "Active"}),
        "completed_students": await db.students.count_documents({"status": "Completed"}),
    }


# ------------------------------------------------------------------ STUDENTS
ACTIVE_ENROLL = ["Registered", "Active", "On Hold"]

def gen_student_id(n): return f"CW-STU-{n:04d}"

async def batch_seats(batch_id):
    b = await db.batches.find_one({"id": batch_id}, {"_id": 0})
    if not b:
        return None
    cap = int(b.get("seats") or 0)
    occ = await db.student_enrollments.count_documents({"batch_id": batch_id, "status": {"$in": ACTIVE_ENROLL}})
    return {"batch": b, "capacity": cap, "occupied": occ, "available": max(cap - occ, 0)}

async def make_enrollment(student_id, course_id, batch_id, joining_date, status):
    c = await db.courses.find_one({"id": course_id}, {"_id": 0}) or {}
    b = await db.batches.find_one({"id": batch_id}, {"_id": 0}) or {}
    doc = {"id": new_id(), "student_id": student_id, "course_id": course_id, "batch_id": batch_id,
           "course_name": c.get("name", ""), "batch_label": f"{c.get('name', 'Batch')} — {b.get('start_date', '')}",
           "enrollment_date": now_iso()[:10], "joining_date": joining_date or now_iso()[:10],
           "status": status, "created_at": now_iso(), "updated_at": now_iso()}
    await db.student_enrollments.insert_one(dict(doc))
    return doc

async def enrich_student(s):
    if not s:
        return s
    s["batch"] = (await db.batches.find_one({"id": s.get("batch_id")}, {"_id": 0})) if s.get("batch_id") else {}
    s["course_info"] = (await db.courses.find_one({"id": s.get("course_id")}, {"_id": 0})) if s.get("course_id") else {}
    s["enrollments"] = await db.student_enrollments.find({"student_id": s["id"]}, {"_id": 0}).sort("created_at", -1).to_list(100)
    return s

@api.get("/admin/students")
async def list_students(admin=Depends(get_current_admin), search: Optional[str] = None,
                        status: Optional[str] = None, course_id: Optional[str] = None,
                        batch_id: Optional[str] = None, mode: Optional[str] = None,
                        page: int = 1, page_size: int = 20):
    q = {}
    if status and status != "All": q["status"] = status
    if course_id and course_id != "All": q["course_id"] = course_id
    if batch_id and batch_id != "All": q["batch_id"] = batch_id
    if mode and mode != "All": q["training_mode"] = mode
    if search:
        rx = {"$regex": re.escape(search), "$options": "i"}
        q["$or"] = [{"full_name": rx}, {"email": rx}, {"mobile": rx}, {"student_id": rx}]
    total = await db.students.count_documents(q)
    items = await db.students.find(q, {"_id": 0}).sort("created_at", -1) \
        .skip((page - 1) * page_size).limit(page_size).to_list(page_size)
    for s in items:
        b = await db.batches.find_one({"id": s.get("batch_id")}, {"_id": 0}) if s.get("batch_id") else None
        c = await db.courses.find_one({"id": s.get("course_id")}, {"_id": 0, "name": 1}) if s.get("course_id") else None
        s["course_name"] = (c or {}).get("name", "")
        s["batch_timing"] = f"{(b or {}).get('days', '')} {(b or {}).get('time', '')}".strip() if b else ""
        s["batch_label"] = f"{(c or {}).get('name', '')} — {(b or {}).get('start_date', '')}" if b else ""
        s["trainer"] = s.get("trainer") or (b or {}).get("trainer", "")
    return {"items": items, "total": total, "page": page, "page_size": page_size}

@api.get("/admin/students/{sid}")
async def get_student(sid: str, admin=Depends(get_current_admin)):
    s = await db.students.find_one({"id": sid}, {"_id": 0})
    if not s:
        raise HTTPException(404, "Student not found")
    return await enrich_student(s)

@api.post("/admin/students")
async def create_student(body: Dict[str, Any], admin=Depends(get_current_admin)):
    override = bool(body.pop("override_capacity", False))
    course_id, batch_id = body.get("course_id"), body.get("batch_id")
    if batch_id:
        info = await batch_seats(batch_id)
        if info and not override and info["available"] <= 0:
            raise HTTPException(400, "Batch is full. Please select another batch.")
    n = await db.students.count_documents({}) + 1
    sid = new_id()
    body["id"] = sid
    body["student_id"] = body.get("student_id") or gen_student_id(n)
    body.setdefault("status", "Registered")
    body.setdefault("payment_status", "Pending")
    body["created_at"] = now_iso(); body["updated_at"] = now_iso()
    await db.students.insert_one(dict(body))
    if course_id and batch_id:
        await make_enrollment(sid, course_id, batch_id, body.get("joining_date", ""), body.get("status", "Registered"))
    return await enrich_student(await db.students.find_one({"id": sid}, {"_id": 0}))

@api.put("/admin/students/{sid}")
async def update_student(sid: str, body: Dict[str, Any], admin=Depends(get_current_admin)):
    body.pop("id", None); body.pop("_id", None)
    body.pop("enrollments", None); body.pop("batch", None); body.pop("course_info", None)
    body["updated_at"] = now_iso()
    r = await db.students.update_one({"id": sid}, {"$set": body})
    if r.matched_count == 0:
        raise HTTPException(404, "Student not found")
    if body.get("status"):
        latest = await db.student_enrollments.find_one({"student_id": sid}, {"_id": 0}, sort=[("created_at", -1)])
        if latest:
            await db.student_enrollments.update_one({"id": latest["id"]}, {"$set": {"status": body["status"], "updated_at": now_iso()}})
    return await enrich_student(await db.students.find_one({"id": sid}, {"_id": 0}))

@api.post("/admin/students/{sid}/enroll")
async def add_enrollment(sid: str, body: Dict[str, Any], admin=Depends(get_current_admin)):
    if not (body.get("course_id") and body.get("batch_id")):
        raise HTTPException(400, "Course and batch are required")
    override = bool(body.pop("override_capacity", False))
    info = await batch_seats(body["batch_id"])
    if info and not override and info["available"] <= 0:
        raise HTTPException(400, "Batch is full. Please select another batch.")
    e = await make_enrollment(sid, body["course_id"], body["batch_id"], body.get("joining_date", ""), body.get("status", "Active"))
    await db.students.update_one({"id": sid}, {"$set": {
        "course_id": body["course_id"], "batch_id": body["batch_id"], "updated_at": now_iso()}})
    return e

@api.delete("/admin/students/{sid}")
async def delete_student(sid: str, admin=Depends(get_current_admin)):
    await db.students.delete_one({"id": sid})
    await db.student_enrollments.delete_many({"student_id": sid})
    return {"ok": True}

@api.post("/admin/enquiries/{eid}/convert")
async def convert_enquiry(eid: str, body: Dict[str, Any], admin=Depends(get_current_admin)):
    enq = await db.enquiries.find_one({"id": eid}, {"_id": 0})
    if not enq:
        raise HTTPException(404, "Enquiry not found")
    if await db.students.find_one({"enquiry_id": eid}):
        raise HTTPException(400, "This enquiry has already been converted to a student")
    course_id = body.get("course_id") or enq.get("course_id")
    batch_id = body.get("batch_id") or enq.get("batch_id")
    override = bool(body.get("override_capacity", False))
    if batch_id:
        info = await batch_seats(batch_id)
        if info and not override and info["available"] <= 0:
            raise HTTPException(400, "Batch is full. Please select another batch.")
    n = await db.students.count_documents({}) + 1
    sid = new_id()
    doc = {"id": sid, "student_id": gen_student_id(n), "enquiry_id": eid,
           "full_name": enq.get("name", ""), "email": enq.get("email", ""), "mobile": enq.get("mobile", ""),
           "whatsapp": enq.get("whatsapp", ""), "city": enq.get("city", ""),
           "course_id": course_id, "batch_id": batch_id,
           "training_mode": body.get("training_mode") or enq.get("mode", ""),
           "trainer": body.get("trainer", ""), "photo_url": body.get("photo_url", ""),
           "joining_date": body.get("joining_date") or now_iso()[:10],
           "status": "Registered", "payment_status": body.get("payment_status", "Pending"),
           "notes": body.get("notes", ""), "created_at": now_iso(), "updated_at": now_iso()}
    await db.students.insert_one(dict(doc))
    if course_id and batch_id:
        await make_enrollment(sid, course_id, batch_id, doc["joining_date"], "Registered")
    await db.enquiries.update_one({"id": eid}, {"$set": {"status": "Converted", "updated_at": now_iso()}})
    return await enrich_student(await db.students.find_one({"id": sid}, {"_id": 0}))

@api.get("/admin/batches/{batch_id}/students")
async def batch_students(batch_id: str, admin=Depends(get_current_admin)):
    info = await batch_seats(batch_id) or {"capacity": 0, "occupied": 0, "available": 0, "batch": {}}
    studs = await db.students.find({"batch_id": batch_id}, {"_id": 0}).sort("created_at", -1).to_list(500)
    return {"capacity": info["capacity"], "occupied": info["occupied"], "available": info["available"],
            "batch": info.get("batch", {}), "students": studs}

@api.get("/admin/students-stats")
async def students_stats(admin=Depends(get_current_admin)):
    studs = await db.students.find({}, {"_id": 0}).to_list(10000)
    courses = {c["id"]: c["name"] for c in await db.courses.find({}, {"_id": 0, "id": 1, "name": 1}).to_list(500)}
    by_status, by_course, by_month = {}, {}, {}
    for s in studs:
        st = s.get("status", "Registered"); by_status[st] = by_status.get(st, 0) + 1
        cn = courses.get(s.get("course_id"), "Unassigned"); by_course[cn] = by_course.get(cn, 0) + 1
        m = (s.get("created_at") or "")[:7]
        if m: by_month[m] = by_month.get(m, 0) + 1
    return {
        "total": len(studs), "registered": by_status.get("Registered", 0), "active": by_status.get("Active", 0),
        "completed": by_status.get("Completed", 0), "on_hold": by_status.get("On Hold", 0),
        "dropped": by_status.get("Dropped", 0) + by_status.get("Cancelled", 0),
        "by_status": [{"name": k, "value": v} for k, v in by_status.items()],
        "by_course": sorted([{"name": k, "value": v} for k, v in by_course.items()], key=lambda x: -x["value"])[:6],
        "by_month": [{"name": k, "value": v} for k, v in sorted(by_month.items())][-6:],
    }

@api.get("/admin/students-export")
async def export_students(admin=Depends(get_current_admin)):
    studs = await db.students.find({}, {"_id": 0}).sort("created_at", -1).to_list(10000)
    courses = {c["id"]: c["name"] for c in await db.courses.find({}, {"_id": 0, "id": 1, "name": 1}).to_list(500)}
    buf = io.StringIO(); w = csv.writer(buf)
    w.writerow(["Student ID", "Name", "Mobile", "Email", "WhatsApp", "Course", "Batch",
                "Trainer", "Mode", "Joining Date", "Status", "Registration Date"])
    for s in studs:
        b = await db.batches.find_one({"id": s.get("batch_id")}, {"_id": 0}) if s.get("batch_id") else {}
        w.writerow([s.get("student_id"), s.get("full_name"), s.get("mobile"), s.get("email"),
                    s.get("whatsapp"), courses.get(s.get("course_id"), ""),
                    (b or {}).get("start_date", ""), s.get("trainer") or (b or {}).get("trainer", ""),
                    s.get("training_mode"), s.get("joining_date"), s.get("status"), (s.get("created_at") or "")[:10]])
    buf.seek(0)
    return StreamingResponse(iter([buf.getvalue()]), media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=students.csv"})


# ================================================================ STUDENT PORTAL
LEARNING_TYPES = ["Video", "Audio", "Audio Overview", "Video Overview", "Notes", "PDF", "Slide Deck", "Mind Map", "Reports", "Flashcards", "Quiz", "Infographic", "Data Table", "Document", "External Resource"]
STUDENT_EDITABLE = {"first_name", "last_name", "mobile", "dob", "gender", "address", "city", "state", "country", "pincode", "photo_url"}

def create_student_token(sid, email):
    payload = {"sub": sid, "email": email or "", "role": "student",
               "exp": datetime.now(timezone.utc) + timedelta(days=7), "type": "access"}
    return jwt.encode(payload, JWT_SECRET, algorithm=JWT_ALG)

async def get_current_student(request: Request) -> dict:
    auth = request.headers.get("Authorization", "")
    token = auth[7:] if auth.startswith("Bearer ") else request.cookies.get("student_token")
    if not token:
        raise HTTPException(401, "Not authenticated")
    try:
        payload = jwt.decode(token, JWT_SECRET, algorithms=[JWT_ALG])
    except jwt.ExpiredSignatureError:
        raise HTTPException(401, "Token expired")
    except jwt.InvalidTokenError:
        raise HTTPException(401, "Invalid token")
    if payload.get("role") != "student":
        raise HTTPException(403, "Not a student session")
    s = await db.students.find_one({"id": payload["sub"]}, {"_id": 0, "password_hash": 0})
    if not s:
        raise HTTPException(401, "Student not found")
    return s

async def log_activity(student_id, action, detail=""):
    await db.student_activity.insert_one({"id": new_id(), "student_id": student_id,
        "action": action, "detail": detail, "created_at": now_iso()})

async def fee_summary(s):
    course = await db.courses.find_one({"id": s.get("course_id")}, {"_id": 0}) if s.get("course_id") else {}
    total = float((course or {}).get("discounted_fee") or (course or {}).get("fee") or s.get("total_fee") or 0)
    paid = float(s.get("paid_amount") or 0)
    remaining = max(total - paid, 0)
    status = "PAID" if total > 0 and remaining <= 0 else ("PARTIAL" if paid > 0 else "UNPAID")
    return {"total_fee": total, "paid": paid, "remaining": remaining, "payment_status": status}

def entitlement_active(e):
    if not e or e.get("status") != "GRANTED":
        return False
    today = now_iso()[:10]
    if e.get("start_date") and e["start_date"] > today:
        return False
    if e.get("expiry_date") and e["expiry_date"] < today:
        return False
    return True

class StudentLoginIn(BaseModel):
    identifier: str
    password: str

@api.post("/student/login")
async def student_login(body: StudentLoginIn):
    ident = body.identifier.strip()
    s = await db.students.find_one({"$or": [{"email": ident.lower()}, {"student_id": ident}, {"student_id": ident.upper()}]})
    if not s or not s.get("password_hash") or not verify_password(body.password, s["password_hash"]):
        raise HTTPException(401, "Invalid credentials. Contact the institute if you haven't set a password.")
    token = create_student_token(s["id"], s.get("email"))
    await log_activity(s["id"], "Login", "Student logged in")
    return {"token": token, "student": {"id": s["id"], "student_id": s.get("student_id"),
            "name": s.get("full_name"), "email": s.get("email")}}

@api.get("/student/me")
async def student_me(s=Depends(get_current_student)):
    return {"id": s["id"], "student_id": s.get("student_id"), "name": s.get("full_name"), "email": s.get("email"), "photo_url": s.get("photo_url")}

@api.get("/student/profile")
async def student_get_profile(s=Depends(get_current_student)):
    return await db.students.find_one({"id": s["id"]}, {"_id": 0, "password_hash": 0})

@api.put("/student/profile")
async def student_update_profile(body: Dict[str, Any], s=Depends(get_current_student)):
    patch = {k: v for k, v in body.items() if k in STUDENT_EDITABLE}
    patch["updated_at"] = now_iso()
    await db.students.update_one({"id": s["id"]}, {"$set": patch})
    await log_activity(s["id"], "Profile Updated")
    return await db.students.find_one({"id": s["id"]}, {"_id": 0, "password_hash": 0})

@api.get("/student/dashboard")
async def student_dashboard(s=Depends(get_current_student)):
    fee = await fee_summary(s)
    course = await db.courses.find_one({"id": s.get("course_id")}, {"_id": 0}) if s.get("course_id") else {}
    batch = await db.batches.find_one({"id": s.get("batch_id")}, {"_id": 0}) if s.get("batch_id") else {}
    resources = await db.learning_resources.find({"course_id": s.get("course_id"), "active": True}, {"_id": 0}).to_list(1000)
    ents = {e["resource_id"]: e for e in await db.entitlements.find({"student_id": s["id"]}, {"_id": 0}).to_list(2000)}
    granted = [r for r in resources if entitlement_active(ents.get(r["id"]))]

    def ct(pred):
        tot = [r for r in resources if pred(r.get("resource_type"))]
        g = [r for r in tot if entitlement_active(ents.get(r["id"]))]
        return {"granted": len(g), "total": len(tot)}

    completed = await db.student_progress.count_documents({"student_id": s["id"], "completed": True})
    base = len(granted) or 1
    progress = min(round((completed / base) * 100), 100) if granted else 0
    quiz_res = [r for r in resources if r.get("resource_type") == "Quiz"]
    quiz_unlocked = any(entitlement_active(ents.get(r["id"])) for r in quiz_res)
    cert_count = await db.certificates.count_documents({"student_id": s["id"]})
    return {
        "student": {k: s.get(k) for k in ["id", "student_id", "full_name", "email", "mobile", "photo_url", "status"]},
        "course": {"name": (course or {}).get("name"), "id": (course or {}).get("id"), "duration": (course or {}).get("duration")},
        "batch": {"label": (f"{(course or {}).get('name', '')} — {(batch or {}).get('start_date', '')}" if batch else ""),
                  "start_date": (batch or {}).get("start_date"), "end_date": (batch or {}).get("end_date"),
                  "trainer": s.get("trainer") or (batch or {}).get("trainer")},
        "fee": fee,
        "access": {
            "videos": ct(lambda t: t in ("Video", "Video Overview")),
            "notes": ct(lambda t: t in ("Notes", "PDF", "Document", "Slide Deck")),
            "audio": ct(lambda t: t in ("Audio", "Audio Overview")),
            "quiz": ct(lambda t: t == "Quiz"),
            "granted_total": len(granted), "resources_total": len(resources),
        },
        "progress": progress,
        "final_quiz": {"unlocked": quiz_unlocked, "exists": len(quiz_res) > 0},
        "certificates": cert_count,
    }

@api.get("/student/resources")
async def student_resources(s=Depends(get_current_student)):
    resources = await db.learning_resources.find({"course_id": s.get("course_id"), "active": True}, {"_id": 0}).to_list(1000)
    resources.sort(key=lambda x: (x.get("module", ""), x.get("sort_order", 0), x.get("created_at", "")))
    ents = {e["resource_id"]: e for e in await db.entitlements.find({"student_id": s["id"]}, {"_id": 0}).to_list(2000)}
    prog = {p["resource_id"] for p in await db.student_progress.find({"student_id": s["id"], "completed": True}, {"_id": 0}).to_list(2000)}
    out = []
    for r in resources:
        e = ents.get(r["id"])
        active = entitlement_active(e)
        reason = None
        if not active:
            if e and e.get("status") == "GRANTED" and e.get("expiry_date") and e["expiry_date"] < now_iso()[:10]:
                reason = "Access expired"
            elif e and e.get("status") == "REVOKED":
                reason = "Access not granted"
            else:
                reason = "Please complete payment for access"
        out.append({
            "id": r["id"], "title": r.get("title"), "resource_type": r.get("resource_type"),
            "module": r.get("module"), "topic": r.get("topic"), "description": r.get("description"),
            "download_allowed": bool(r.get("download_allowed")), "completed": r["id"] in prog,
            "locked": not active, "lock_reason": reason,
            "access": {"start_date": (e or {}).get("start_date"), "expiry_date": (e or {}).get("expiry_date")},
        })
    return out

@api.get("/student/resources/{rid}/access")
async def student_resource_access(rid: str, s=Depends(get_current_student)):
    r = await db.learning_resources.find_one({"id": rid}, {"_id": 0})
    if not r:
        raise HTTPException(404, "Resource not found")
    e = await db.entitlements.find_one({"student_id": s["id"], "resource_id": rid}, {"_id": 0})
    if not entitlement_active(e):
        raise HTTPException(403, "You do not have access to this resource.")
    await log_activity(s["id"], "Resource Opened", r.get("title", ""))
    return {"id": r["id"], "title": r.get("title"), "resource_type": r.get("resource_type"),
            "content_url": r.get("content_url") or r.get("external_url") or "",
            "download_allowed": bool(r.get("download_allowed")), "body": r.get("body", "")}

@api.post("/student/resources/{rid}/complete")
async def student_resource_complete(rid: str, s=Depends(get_current_student)):
    e = await db.entitlements.find_one({"student_id": s["id"], "resource_id": rid}, {"_id": 0})
    if not entitlement_active(e):
        raise HTTPException(403, "No access")
    await db.student_progress.update_one({"student_id": s["id"], "resource_id": rid},
        {"$set": {"completed": True, "updated_at": now_iso()}, "$setOnInsert": {"id": new_id()}}, upsert=True)
    await log_activity(s["id"], "Resource Completed", rid)
    return {"ok": True}

@api.post("/student/documents")
async def student_upload_doc(file: UploadFile = File(...), doc_type: str = Query("Other"), s=Depends(get_current_student)):
    data = await file.read()
    if len(data) > 10 * 1024 * 1024:
        raise HTTPException(400, "File too large (max 10MB)")
    mid = new_id()
    await db.media_files.insert_one({"id": mid, "content_type": file.content_type or "application/octet-stream",
        "data": base64.b64encode(data).decode(), "private": True, "owner_student": s["id"], "created_at": now_iso()})
    doc = {"id": new_id(), "student_id": s["id"], "doc_type": doc_type, "file_id": mid,
           "filename": file.filename, "created_at": now_iso()}
    await db.student_documents.insert_one(dict(doc))
    await log_activity(s["id"], "Document Uploaded", doc_type)
    doc.pop("_id", None)
    return doc

@api.get("/student/documents")
async def student_list_docs(s=Depends(get_current_student)):
    return await db.student_documents.find({"student_id": s["id"]}, {"_id": 0}).sort("created_at", -1).to_list(200)

@api.get("/student/documents/{doc_id}/file")
async def student_doc_file(doc_id: str, s=Depends(get_current_student)):
    d = await db.student_documents.find_one({"id": doc_id, "student_id": s["id"]}, {"_id": 0})
    if not d:
        raise HTTPException(404, "Document not found")
    m = await db.media_files.find_one({"id": d["file_id"]}, {"_id": 0})
    if not m:
        raise HTTPException(404, "File missing")
    await log_activity(s["id"], "Download", d.get("doc_type", ""))
    return Response(content=base64.b64decode(m["data"]), media_type=m["content_type"],
        headers={"Content-Disposition": f'inline; filename="{d.get("filename", "document")}"'})

@api.delete("/student/documents/{doc_id}")
async def student_delete_doc(doc_id: str, s=Depends(get_current_student)):
    d = await db.student_documents.find_one({"id": doc_id, "student_id": s["id"]}, {"_id": 0})
    if not d:
        raise HTTPException(404, "Document not found")
    await db.media_files.delete_one({"id": d["file_id"]})
    await db.student_documents.delete_one({"id": doc_id})
    return {"ok": True}

@api.get("/student/activity")
async def student_activity(s=Depends(get_current_student)):
    return await db.student_activity.find({"student_id": s["id"]}, {"_id": 0}).sort("created_at", -1).to_list(200)

@api.get("/student/certificates")
async def student_certificates(s=Depends(get_current_student)):
    return await db.certificates.find({"student_id": s["id"]}, {"_id": 0}).sort("created_at", -1).to_list(100)

# ---- admin: student password, learning access, entitlements, documents
@api.post("/admin/students/{sid}/set-password")
async def admin_set_student_password(sid: str, body: Dict[str, Any], admin=Depends(get_current_admin)):
    pw = body.get("password")
    if not pw or len(pw) < 6:
        raise HTTPException(400, "Password must be at least 6 characters")
    r = await db.students.update_one({"id": sid}, {"$set": {"password_hash": hash_password(pw), "updated_at": now_iso()}})
    if r.matched_count == 0:
        raise HTTPException(404, "Student not found")
    return {"ok": True}

@api.get("/admin/students/{sid}/learning")
async def admin_student_learning(sid: str, admin=Depends(get_current_admin)):
    s = await db.students.find_one({"id": sid}, {"_id": 0, "password_hash": 0})
    if not s:
        raise HTTPException(404, "Student not found")
    resources = await db.learning_resources.find({"course_id": s.get("course_id")}, {"_id": 0}).to_list(1000)
    resources.sort(key=lambda x: (x.get("module", ""), x.get("sort_order", 0)))
    ents = {e["resource_id"]: e for e in await db.entitlements.find({"student_id": sid}, {"_id": 0}).to_list(2000)}
    for r in resources:
        e = ents.get(r["id"])
        r["access_status"] = (e or {}).get("status", "NONE")
        r["start_date"] = (e or {}).get("start_date")
        r["expiry_date"] = (e or {}).get("expiry_date")
    return {"student": s, "fee": await fee_summary(s), "resources": resources,
            "has_password": bool((await db.students.find_one({"id": sid}, {"password_hash": 1})).get("password_hash"))}

@api.get("/admin/students/{sid}/documents")
async def admin_student_documents(sid: str, admin=Depends(get_current_admin)):
    return await db.student_documents.find({"student_id": sid}, {"_id": 0}).sort("created_at", -1).to_list(200)

@api.get("/admin/documents/{doc_id}/file")
async def admin_doc_file(doc_id: str, admin=Depends(get_current_admin)):
    d = await db.student_documents.find_one({"id": doc_id}, {"_id": 0})
    if not d:
        raise HTTPException(404, "Not found")
    m = await db.media_files.find_one({"id": d["file_id"]}, {"_id": 0})
    if not m:
        raise HTTPException(404, "File missing")
    return Response(content=base64.b64decode(m["data"]), media_type=m["content_type"],
        headers={"Content-Disposition": f'inline; filename="{d.get("filename", "document")}"'})

@api.get("/admin/students/{sid}/activity")
async def admin_student_activity(sid: str, admin=Depends(get_current_admin)):
    return await db.student_activity.find({"student_id": sid}, {"_id": 0}).sort("created_at", -1).to_list(300)

class EntitlementIn(BaseModel):
    student_ids: List[str]
    resource_ids: List[str]
    start_date: Optional[str] = None
    expiry_date: Optional[str] = None

@api.post("/admin/entitlements/grant")
async def admin_grant(body: EntitlementIn, admin=Depends(get_current_admin)):
    n = 0
    for sid in body.student_ids:
        for rid in body.resource_ids:
            await db.entitlements.update_one({"student_id": sid, "resource_id": rid},
                {"$set": {"student_id": sid, "resource_id": rid, "status": "GRANTED",
                          "start_date": body.start_date or now_iso()[:10], "expiry_date": body.expiry_date or "",
                          "granted_by": admin.get("email"), "granted_at": now_iso(), "revoked_at": "", "updated_at": now_iso()},
                 "$setOnInsert": {"id": new_id()}}, upsert=True)
            await log_activity(sid, "Access Granted", rid)
            n += 1
    return {"ok": True, "count": n}

@api.post("/admin/entitlements/revoke")
async def admin_revoke(body: EntitlementIn, admin=Depends(get_current_admin)):
    n = 0
    for sid in body.student_ids:
        for rid in body.resource_ids:
            await db.entitlements.update_one({"student_id": sid, "resource_id": rid},
                {"$set": {"status": "REVOKED", "revoked_at": now_iso(), "updated_at": now_iso()}})
            await log_activity(sid, "Access Revoked", rid)
            n += 1
    return {"ok": True, "count": n}

@app.on_event("startup")
async def seed_student_portal():
    try:
        await db.entitlements.create_index([("student_id", 1), ("resource_id", 1)], unique=True)
        await db.students.create_index("student_id")
        await db.student_activity.create_index("student_id")
        if await db.students.count_documents({}) == 0:
            course = await db.courses.find_one({"published": True}, {"_id": 0})
            if course:
                batch = await db.batches.find_one({}, {"_id": 0})
                sid = new_id()
                await db.students.insert_one({"id": sid, "student_id": "CW-2026-0001",
                    "full_name": "Deepak Harale", "first_name": "Deepak", "last_name": "Harale",
                    "email": "student@cloudwavetechnologies.com", "mobile": "9000000000",
                    "course_id": course["id"], "batch_id": (batch or {}).get("id", ""),
                    "trainer": (batch or {}).get("trainer", "CloudWave Trainer"),
                    "status": "Active", "payment_status": "Partial", "paid_amount": 20000,
                    "total_fee": course.get("discounted_fee") or course.get("fee") or 50000,
                    "password_hash": hash_password("Student@1234"),
                    "created_at": now_iso(), "updated_at": now_iso()})
                if batch:
                    await make_enrollment(sid, course["id"], batch["id"], now_iso()[:10], "Active")
                defs = [("Course Introduction", "Video Overview", 1, True), ("Module 1 — Core Concepts", "Video", 1, True),
                        ("Module 1 Notes", "Notes", 1, True), ("Module 2 — Advanced", "Video", 2, False),
                        ("Advanced Notes", "PDF", 2, False), ("Final Course Test", "Quiz", 3, False)]
                for i, (title, rtype, mod, grant) in enumerate(defs):
                    rid = new_id()
                    await db.learning_resources.insert_one({"id": rid, "title": title, "resource_type": rtype,
                        "course_id": course["id"], "batch_id": (batch or {}).get("id", ""), "module": f"Module {mod}",
                        "content_url": "", "external_url": "", "download_allowed": rtype in ("Notes", "PDF"),
                        "body": "Sample learning content for " + title + ".", "sort_order": i, "active": True,
                        "created_at": now_iso(), "updated_at": now_iso()})
                    if grant:
                        await db.entitlements.insert_one({"id": new_id(), "student_id": sid, "resource_id": rid,
                            "status": "GRANTED", "start_date": now_iso()[:10], "expiry_date": "",
                            "granted_by": "system", "granted_at": now_iso(), "updated_at": now_iso()})
    except Exception as e:
        logger.error(f"student portal seed failed: {e}")


# ------------------------------------------------------------------ PAYMENTS (provider-agnostic)
# Initial credential seed values (used only to seed the DB provider once; runtime reads from DB).
ENV_CLOUDPAY = {
    "base_url": os.environ.get("CLOUDPAY_BASE_URL", "").strip(),
    "api_key": os.environ.get("CLOUDPAY_API_KEY", "").strip(),
    "secret": os.environ.get("CLOUDPAY_SECRET", "").strip(),
}
SECRET_FIELDS = ("api_key", "secret")

def course_price(course, currency):
    if currency == "USD":
        return float(course.get("price_usd") or 0)
    return float(course.get("discounted_fee") or course.get("fee") or 0)

def provider_configured(prov):
    prov = prov or {}
    return bool(prov.get("base_url") and prov.get("api_key") and prov.get("secret"))

def provider_is_sandbox(prov):
    # Sandbox (mock) when explicitly in sandbox mode OR when Live credentials are incomplete.
    prov = prov or {}
    return prov.get("mode", "sandbox") != "live" or not provider_configured(prov)

def mask_provider(p):
    out = dict(p)
    for f in SECRET_FIELDS:
        val = p.get(f) or ""
        out[f + "_set"] = bool(val)
        out[f] = ("\u2022\u2022\u2022\u2022" + val[-4:]) if len(val) >= 4 else ("\u2022\u2022\u2022\u2022" if val else "")
    out["configured"] = provider_configured(p)
    out["sandbox"] = provider_is_sandbox(p)
    return out

async def get_provider(pid):
    return await db.payment_providers.find_one({"id": pid}, {"_id": 0})

async def active_provider():
    provs = await db.payment_providers.find({"enabled": True}, {"_id": 0}).to_list(50)
    provs.sort(key=lambda x: x.get("priority", 0), reverse=True)
    return provs[0] if provs else None

async def cloudpay_create(prov, order):
    # Sandbox/mock flow: hand the student off to our hosted mock checkout page.
    if provider_is_sandbox(prov):
        return {"status": "created", "gateway_ref": "SANDBOX-" + order["order_ref"],
                "checkout_url": f"/payment/checkout?ref={order['order_ref']}", "sandbox": True}
    try:
        async with httpx.AsyncClient(timeout=20) as c:
            r = await c.post(f"{prov['base_url'].rstrip('/')}/orders",
                headers={"Authorization": f"Bearer {prov['api_key']}", "X-Mode": prov.get("mode", "live")},
                json={"amount": order["amount"], "currency": order["currency"], "reference": order["order_ref"],
                      "customer": order.get("customer", {})})
        r.raise_for_status()
        d = r.json()
        return {"status": "created", "gateway_ref": d.get("id"), "checkout_url": d.get("checkout_url")}
    except Exception as e:
        logger.error(f"CloudPay create failed: {e}")
        return {"status": "error", "message": "Unable to start payment. Please try again."}

async def cloudpay_verify(prov, order, sandbox_result="success"):
    if provider_is_sandbox(prov):
        return {"paid": sandbox_result != "fail"}
    ref = order.get("gateway_ref")
    if not ref:
        return {"paid": False}
    try:
        async with httpx.AsyncClient(timeout=20) as c:
            r = await c.get(f"{prov['base_url'].rstrip('/')}/orders/{ref}",
                headers={"Authorization": f"Bearer {prov['api_key']}", "X-Mode": prov.get("mode", "live")})
        r.raise_for_status()
        return {"paid": r.json().get("status") in ("paid", "success", "captured")}
    except Exception as e:
        logger.error(f"CloudPay verify failed: {e}")
        return {"paid": False}

async def provider_create_order(prov, order):
    if prov and prov["id"] == "cloudpay":
        return await cloudpay_create(prov, order)
    return {"status": "error", "message": f"Provider '{(prov or {}).get('id')}' is not enabled."}

async def provider_verify(prov, order, sandbox_result="success"):
    if prov and prov["id"] == "cloudpay":
        return await cloudpay_verify(prov, order, sandbox_result)
    return {"paid": False}

async def confirm_enrollment(order):
    if order.get("student_id"):
        return order["student_id"]
    n = await db.students.count_documents({}) + 1
    sid = new_id()
    cust = order.get("customer", {})
    await db.students.insert_one({
        "id": sid, "student_id": gen_student_id(n), "full_name": cust.get("name", ""),
        "email": cust.get("email", ""), "mobile": cust.get("mobile", ""),
        "course_id": order.get("course_id"), "batch_id": order.get("batch_id"),
        "status": "Registered", "payment_status": "Paid", "training_mode": "",
        "joining_date": now_iso()[:10], "created_at": now_iso(), "updated_at": now_iso()})
    if order.get("course_id") and order.get("batch_id"):
        await make_enrollment(sid, order["course_id"], order["batch_id"], now_iso()[:10], "Registered")
    await db.orders.update_one({"id": order["id"]}, {"$set": {"student_id": sid, "updated_at": now_iso()}})
    return sid

class OrderIn(BaseModel):
    course_id: str
    batch_id: Optional[str] = ""
    currency: str = "INR"
    name: str
    email: EmailStr
    mobile: str

@api.get("/payment/config")
async def payment_config():
    p = await active_provider()
    return {"provider": p["id"] if p else None, "name": (p or {}).get("name"),
            "currencies": (p or {}).get("currencies", ["INR", "USD"]),
            "sandbox": provider_is_sandbox(p) if p else True,
            "configured": provider_configured(p) if p else False}

@api.post("/payment/create-order")
async def create_order(body: OrderIn):
    course = await db.courses.find_one({"id": body.course_id, "published": True}, {"_id": 0})
    if not course:
        course = await db.courses.find_one({"slug": body.course_id}, {"_id": 0})
    if not course:
        raise HTTPException(404, "Course not found")
    if body.currency not in ("INR", "USD"):
        raise HTTPException(400, "Unsupported currency")
    amount = course_price(course, body.currency)
    if amount <= 0:
        raise HTTPException(400, f"This course has no {body.currency} price configured.")
    prov = await active_provider()
    if not prov:
        raise HTTPException(503, "No payment provider is enabled. Please contact the institute.")
    if body.batch_id:
        info = await batch_seats(body.batch_id)
        if info and info["available"] <= 0:
            raise HTTPException(400, "Batch is full. Please select another batch.")
    count = await db.orders.count_documents({}) + 1
    order = {"id": new_id(), "order_ref": f"CWO{datetime.now().strftime('%y%m')}{count:05d}",
             "course_id": course["id"], "course_name": course["name"], "batch_id": body.batch_id or "",
             "amount": amount, "currency": body.currency, "provider": prov["id"], "status": "pending",
             "gateway_ref": "", "student_id": "", "sandbox": provider_is_sandbox(prov),
             "customer": {"name": body.name, "email": body.email, "mobile": body.mobile},
             "created_at": now_iso(), "updated_at": now_iso()}
    res = await provider_create_order(prov, order)
    order["gateway_ref"] = res.get("gateway_ref", "")
    await db.orders.insert_one(dict(order))
    return {"order_ref": order["order_ref"], "provider": prov["id"], "amount": amount, "currency": body.currency,
            "status": res.get("status"), "checkout_url": res.get("checkout_url"),
            "sandbox": order["sandbox"], "message": res.get("message")}

@api.get("/payment/order/{order_ref}")
async def get_order_public(order_ref: str):
    o = await db.orders.find_one({"order_ref": order_ref}, {"_id": 0})
    if not o:
        raise HTTPException(404, "Order not found")
    return {"order_ref": o["order_ref"], "course_name": o.get("course_name"), "amount": o.get("amount"),
            "currency": o.get("currency"), "status": o.get("status"), "provider": o.get("provider"),
            "sandbox": bool(o.get("sandbox")), "customer_name": (o.get("customer") or {}).get("name", "")}

@api.post("/payment/verify")
async def verify_payment(body: Dict[str, Any]):
    order = await db.orders.find_one({"order_ref": body.get("order_ref")}, {"_id": 0})
    if not order:
        raise HTTPException(404, "Order not found")
    prov = await get_provider(order["provider"])
    sandbox_result = body.get("sandbox_result", "success")
    v = await provider_verify(prov, order, sandbox_result)
    if v.get("paid"):
        await db.orders.update_one({"id": order["id"]}, {"$set": {"status": "success", "updated_at": now_iso()}})
        sid = await confirm_enrollment(order)
        return {"status": "success", "student_id": sid}
    if sandbox_result == "fail":
        await db.orders.update_one({"id": order["id"]}, {"$set": {"status": "failed", "updated_at": now_iso()}})
        return {"status": "failed"}
    return {"status": order.get("status", "pending")}

@api.post("/payment/webhook/{provider}")
async def payment_webhook(provider: str, body: Dict[str, Any]):
    ref = body.get("reference") or body.get("order_ref")
    order = await db.orders.find_one({"order_ref": ref}, {"_id": 0})
    if not order:
        return {"ok": True}
    if body.get("id"):
        order["gateway_ref"] = order.get("gateway_ref") or body.get("id")
    prov = await get_provider(order["provider"])
    v = await provider_verify(prov, order)
    if v.get("paid"):
        await db.orders.update_one({"id": order["id"]}, {"$set": {"status": "success", "updated_at": now_iso()}})
        await confirm_enrollment(order)
    return {"ok": True}

@api.get("/admin/payment-providers")
async def list_providers(admin=Depends(get_current_admin)):
    provs = await db.payment_providers.find({}, {"_id": 0}).to_list(50)
    return [mask_provider(p) for p in provs]

@api.put("/admin/payment-providers/{pid}")
async def update_provider(pid: str, body: Dict[str, Any], admin=Depends(get_current_admin)):
    body.pop("id", None); body.pop("configured", None); body.pop("sandbox", None)
    for f in SECRET_FIELDS:
        body.pop(f + "_set", None)
        if f in body:
            v = body.get(f)
            # Keep existing stored secret when the field is blank or still masked.
            if v is None or v == "" or "\u2022" in str(v):
                body.pop(f, None)
    body["updated_at"] = now_iso()
    await db.payment_providers.update_one({"id": pid}, {"$set": body}, upsert=True)
    p = await db.payment_providers.find_one({"id": pid}, {"_id": 0})
    return mask_provider(p)

@api.get("/admin/orders")
async def list_orders(admin=Depends(get_current_admin)):
    return await db.orders.find({}, {"_id": 0}).sort("created_at", -1).to_list(2000)


# ------------------------------------------------------------------ CLAUDE AI ASSISTANT
EMERGENT_LLM_KEY = os.environ.get("EMERGENT_LLM_KEY")

class AssistantIn(BaseModel):
    message: str
    history: Optional[List[Dict[str, Any]]] = []
    session_id: Optional[str] = ""

@api.post("/assistant/chat")
async def assistant_chat(body: AssistantIn):
    if not EMERGENT_LLM_KEY:
        raise HTTPException(503, "AI assistant is not configured")
    from emergentintegrations.llm.chat import LlmChat, UserMessage
    courses = await db.courses.find({"published": True},
        {"_id": 0, "name": 1, "category": 1, "duration": 1, "level": 1, "fee": 1, "discounted_fee": 1, "short_description": 1}).to_list(100)
    settings = await db.website_settings.find_one({"id": "main"}, {"_id": 0}) or {}
    clist = "\n".join(f"- {c['name']} ({c.get('category','')}, {c.get('duration','')}, Fee Rs.{c.get('discounted_fee') or c.get('fee')}): {c.get('short_description','')}" for c in courses)
    sys = (f"You are the friendly AI course assistant for {settings.get('institute_name','CloudWave Technologies')}, an IT training institute. "
           "Help visitors find the right course, explain duration, fees, level and training mode, and warmly encourage them to submit an enquiry or contact the team. "
           "Keep replies concise (2-5 sentences), professional and helpful. Only answer questions about the institute and its IT training. "
           "If asked something unrelated or unknown, politely steer back and suggest contacting the team.\n\n"
           f"Available courses:\n{clist}\n\n"
           f"Contact: phone {settings.get('phone','')}, WhatsApp {settings.get('whatsapp','')}, email {settings.get('email','')}. "
           "To enquire, tell them to click the 'Enquire Now' button or visit the Enquiry page.")
    chat = LlmChat(api_key=EMERGENT_LLM_KEY, session_id=body.session_id or new_id(), system_message=sys).with_model("anthropic", "claude-sonnet-4-6")
    hist = ""
    for h in (body.history or [])[-6:]:
        hist += f"\n{h.get('role','user')}: {h.get('content','')}"
    prompt = (hist + "\nuser: " + body.message).strip() if hist else body.message
    try:
        reply = await chat.send_message(UserMessage(text=prompt))
    except Exception as e:
        logger.error(f"Assistant error: {e}")
        raise HTTPException(500, "Assistant is temporarily unavailable")
    return {"reply": reply if isinstance(reply, str) else str(reply)}


@api.get("/slides")
async def public_slides():
    slides = await pub_list("slides", "published")
    at = (await active_theme()).get("theme")
    if at and (at.get("banner_video") or at.get("banner_image")):
        media = "video" if at.get("banner_video") else "image"
        slides = [{"id": "festival-" + at["id"], "media_type": media,
                   "image_url": at.get("banner_image", ""), "video_url": at.get("banner_video", ""),
                   "title": at.get("name", ""), "description": at.get("announcement_text", ""),
                   "festival": True, "sort_order": -1, "published": True}] + slides
    return slides

@api.get("/admin/batches/{batch_id}/financials")
async def batch_financials(batch_id: str, admin=Depends(get_current_admin)):
    b = await db.batches.find_one({"id": batch_id}, {"_id": 0}) or {}
    course = await db.courses.find_one({"id": b.get("course_id")}, {"_id": 0}) if b.get("course_id") else None
    fee = float((course or {}).get("discounted_fee") or (course or {}).get("fee") or 0)
    students = await db.students.find({"batch_id": batch_id}, {"_id": 0}).to_list(2000)
    orders = await db.orders.find({"batch_id": batch_id, "status": "success"}, {"_id": 0}).to_list(2000)
    expenses = await db.expenses.find({"batch_id": batch_id}, {"_id": 0}).to_list(2000)
    paid = [s for s in students if s.get("payment_status") == "Paid"]
    pending = [s for s in students if s.get("payment_status") != "Paid"]
    online_ids = {o.get("student_id") for o in orders}
    online = sum(float(o.get("amount") or 0) for o in orders)
    offline = sum(float(s.get("paid_amount") or fee) for s in paid if s["id"] not in online_ids)
    total = online + offline
    expected = len(students) * fee
    exp_total = sum(float(e.get("amount") or 0) for e in expenses)
    net = total - exp_total
    margin = (net / total * 100) if total > 0 else 0
    return {"fee": fee, "total_students": len(students), "paid_students": len(paid), "pending_students": len(pending),
            "expected": expected, "online": online, "offline": offline, "total_collection": total,
            "pending_collection": max(expected - total, 0), "expenses_total": exp_total,
            "net_profit": net, "profit_margin": round(margin, 1), "expenses": expenses}

@api.get("/admin/batches/{batch_id}/excel")
async def batch_excel(batch_id: str, admin=Depends(get_current_admin)):
    from openpyxl import Workbook
    fin = await batch_financials(batch_id, admin)
    b = await db.batches.find_one({"id": batch_id}, {"_id": 0}) or {}
    courses = {c["id"]: c["name"] for c in await db.courses.find({}, {"_id": 0, "id": 1, "name": 1}).to_list(500)}
    students = await db.students.find({"batch_id": batch_id}, {"_id": 0}).to_list(2000)
    orders = {o.get("student_id"): o for o in await db.orders.find({"batch_id": batch_id}, {"_id": 0}).to_list(2000)}
    wb = Workbook()
    s1 = wb.active; s1.title = "Students"
    s1.append(["Student ID", "Name", "Course", "Batch", "Payment Type", "Payment Status", "Amount", "Txn/Order Ref", "Address", "Enrollment Date", "Payment Date"])
    for s in students:
        o = orders.get(s["id"])
        s1.append([s.get("student_id"), s.get("full_name"), courses.get(s.get("course_id"), ""), b.get("start_date", ""),
                   "Online" if o else "Offline", s.get("payment_status", ""), (o or {}).get("amount", ""),
                   (o or {}).get("order_ref", ""), s.get("address", ""), (s.get("created_at") or "")[:10], s.get("joining_date", "")])
    s2 = wb.create_sheet("Payment Collection")
    for row in [["Expected Collection", fin["expected"]], ["Online Collection", fin["online"]], ["Offline Collection", fin["offline"]],
                ["Total Successful Collection", fin["total_collection"]], ["Pending Collection", fin["pending_collection"]],
                ["Paid Student Count", fin["paid_students"]], ["Pending Student Count", fin["pending_students"]]]:
        s2.append(row)
    s3 = wb.create_sheet("Expenses")
    s3.append(["Expense ID", "Date", "Category", "Description", "Amount", "Notes"])
    for e in fin["expenses"]:
        s3.append([e.get("id", "")[:8], e.get("expense_date", ""), e.get("category", ""), e.get("description", ""), e.get("amount", ""), e.get("notes", "")])
    s4 = wb.create_sheet("Profit Summary")
    for row in [["Total Collection", fin["total_collection"]], ["Total Expenses", fin["expenses_total"]],
                ["Net Profit", fin["net_profit"]], ["Profit Margin (%)", fin["profit_margin"]]]:
        s4.append(row)
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="batch-{batch_id[:8]}.xlsx"'})


app.include_router(api)

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=os.environ.get('CORS_ORIGINS', '*').split(','),
    allow_methods=["*"],
    allow_headers=["*"],
)


# ------------------------------------------------------------------ SEED
@app.on_event("startup")
async def startup():
    await db.users.create_index("email", unique=True)
    await db.enquiries.create_index("enquiry_id")
    await db.courses.create_index("slug")
    await db.blog_posts.create_index("slug")

    # admins
    for em, pw, nm in [(os.environ["ADMIN_EMAIL"], os.environ["ADMIN_PASSWORD"], "CloudWave Admin"),
                       (os.environ.get("OWNER_EMAIL"), os.environ.get("OWNER_PASSWORD"), "Owner")]:
        if not em:
            continue
        em = em.lower()
        ex = await db.users.find_one({"email": em})
        if not ex:
            await db.users.insert_one({"id": new_id(), "email": em, "name": nm,
                "role": "admin", "password_hash": hash_password(pw), "created_at": now_iso()})
        elif not verify_password(pw, ex["password_hash"]):
            await db.users.update_one({"email": em}, {"$set": {"password_hash": hash_password(pw)}})

    await seed_content()


async def seed_content():
    if not await db.website_settings.find_one({"id": "main"}):
        await db.website_settings.insert_one({
            "id": "main", "institute_name": "CloudWave Technologies",
            "tagline": "Launch Your IT Career in the Cloud",
            "logo_url": "", "phone": "8262030386", "phone_alt": "9921218389",
            "whatsapp": "918262030386", "email": "admin@cloudwavetechnologies.org",
            "address": "CloudWave Technologies Pvt Ltd, Pune, Maharashtra, India",
            "maps_url": "https://www.google.com/maps?q=Pune,Maharashtra",
            "facebook": "https://facebook.com", "instagram": "https://instagram.com",
            "linkedin": "https://linkedin.com", "youtube": "https://youtube.com",
            "twitter": "https://twitter.com",
            "youtube_channel": "https://youtube.com",
            "instagram_profile": "https://instagram.com",
            "working_hours": "Mon - Sat: 9:00 AM - 8:00 PM",
            "hero_title": "Launch Your IT Career with Industry-Ready Training",
            "hero_description": "CloudWave Technologies delivers hands-on, placement-focused training in software development, cloud, testing and data — taught by working industry experts.",
            "footer_text": "CloudWave Technologies Pvt Ltd — building tomorrow's technology professionals.",
            "stat_students": "5000", "stat_courses": "25", "stat_trainers": "40",
            "stat_placement": "92", "stat_experience": "10",
            "theme_mode": "auto",
            "updated_at": now_iso()})

    if await db.courses.count_documents({}) == 0:
        cats = ["Development", "Cloud", "Testing", "Data", "Programming"]
        courses = [
            ("Business Analyst", "Development", "Master requirement analysis, Agile, SQL, Power BI and stakeholder management.", "3 Months", 45000, 32000),
            ("Selenium with Java", "Testing", "Become an automation testing engineer with Selenium, Java, TestNG and frameworks.", "3 Months", 40000, 28000),
            ("AWS Solutions Architect", "Cloud", "Design and deploy scalable cloud infrastructure on Amazon Web Services.", "4 Months", 50000, 38000),
            ("Python Full Stack", "Programming", "Learn Python, Django, REST APIs, React and databases end-to-end.", "5 Months", 55000, 40000),
            ("Java Full Stack", "Programming", "Enterprise Java, Spring Boot, Microservices, React and deployment.", "5 Months", 55000, 42000),
            ("Data Analytics", "Data", "Excel, SQL, Python, Power BI and Tableau for data-driven careers.", "4 Months", 48000, 35000),
            ("Software Testing (Manual + Automation)", "Testing", "Complete QA journey from manual testing to automation frameworks.", "3 Months", 38000, 26000),
            ("DevOps Engineering", "Cloud", "CI/CD, Docker, Kubernetes, Jenkins, Terraform and cloud deployment.", "4 Months", 52000, 40000),
        ]
        levels = ["Beginner", "Intermediate", "Advanced"]
        for i, (name, cat, desc, dur, fee, disc) in enumerate(courses):
            await db.courses.insert_one({
                "id": new_id(), "slug": slugify(name), "name": name, "category": cat,
                "short_description": desc,
                "full_description": desc + " This programme blends live instructor-led sessions, real projects, resume preparation and dedicated placement support to make you job-ready.",
                "image_url": "", "duration": dur, "level": levels[i % 3],
                "mode": "Online / Classroom", "fee": fee, "discounted_fee": disc,
                "certification": "CloudWave Certified " + name.split()[0],
                "prerequisites": "Basic computer knowledge. No prior coding experience required.",
                "learning_outcomes": ["Job-ready practical skills", "Real-time industry projects",
                    "Interview preparation", "Recognised certification"],
                "career_opportunities": ["Software Engineer", "Consultant", "Analyst", "Team Lead"],
                "tools": ["Git", "VS Code", "Jira", "Postman"],
                "syllabus": [
                    {"module": "Module 1: Foundations", "topics": ["Introduction & setup", "Core concepts", "Hands-on basics"]},
                    {"module": "Module 2: Core Skills", "topics": ["Advanced concepts", "Live examples", "Assignments"]},
                    {"module": "Module 3: Projects", "topics": ["Real-world project", "Best practices", "Deployment"]},
                    {"module": "Module 4: Career Prep", "topics": ["Resume building", "Mock interviews", "Placement support"]},
                ],
                "featured": i < 4, "published": True, "sort_order": i,
                "created_at": now_iso(), "updated_at": now_iso()})

        course_docs = await db.courses.find({}, {"_id": 0}).to_list(50)
        cmap = {c["name"]: c["id"] for c in course_docs}

        trainers = [
            ("Rahul Deshmukh", "Lead Automation Architect", "12+ years", ["Selenium", "Java", "TestNG", "CI/CD"]),
            ("Priya Nair", "Cloud & DevOps Mentor", "10+ years", ["AWS", "Docker", "Kubernetes", "Terraform"]),
            ("Amit Verma", "Full Stack Lead", "11+ years", ["Python", "Django", "React", "PostgreSQL"]),
            ("Sneha Kulkarni", "Data Analytics Expert", "9+ years", ["SQL", "Power BI", "Tableau", "Python"]),
            ("Vikram Singh", "Java Enterprise Trainer", "13+ years", ["Spring Boot", "Microservices", "Java", "AWS"]),
            ("Neha Sharma", "Business Analysis Coach", "8+ years", ["Agile", "SQL", "Power BI", "JIRA"]),
        ]
        for i, (nm, desig, exp, skills) in enumerate(trainers):
            await db.trainers.insert_one({
                "id": new_id(), "name": nm, "designation": desig, "experience": exp,
                "skills": skills, "certifications": ["Industry Certified Professional"],
                "bio": f"{nm} is a seasoned professional with {exp} of industry experience, having trained thousands of students into successful IT careers.",
                "photo_url": "", "linkedin": "https://linkedin.com",
                "active": True, "sort_order": i, "created_at": now_iso(), "updated_at": now_iso()})

        import datetime as _dt
        base = _dt.date.today()
        statuses = ["Upcoming", "Filling Fast", "Upcoming", "Started"]
        for i, (name, cid) in enumerate(list(cmap.items())[:5]):
            sd = base + _dt.timedelta(days=7 + i * 10)
            await db.batches.insert_one({
                "id": new_id(), "course_id": cid, "course_name": name,
                "start_date": sd.isoformat(), "end_date": (sd + _dt.timedelta(days=90)).isoformat(),
                "days": "Mon, Wed, Fri", "time": "7:00 PM - 9:00 PM",
                "trainer": trainers[i % len(trainers)][0], "mode": "Online",
                "seats": 30 - i * 3, "location": "Online / Pune Center",
                "status": statuses[i % len(statuses)], "published": True, "sort_order": i,
                "created_at": now_iso(), "updated_at": now_iso()})

        testimonials = [
            ("Aarti Joshi", "Selenium with Java", 5, "The trainers are amazing and the placement support helped me land a QA job within a month of finishing."),
            ("Rohan Mehta", "AWS Solutions Architect", 5, "Hands-on cloud labs and real projects gave me the confidence to clear my AWS certification and interviews."),
            ("Kavya Reddy", "Data Analytics", 5, "From zero coding knowledge to a Data Analyst role — CloudWave changed my career completely."),
            ("Suresh Patil", "Java Full Stack", 4, "Excellent structured curriculum and very supportive mentors. Highly recommend for freshers."),
        ]
        for i, (nm, crs, rt, txt) in enumerate(testimonials):
            await db.testimonials.insert_one({
                "id": new_id(), "name": nm, "course": crs, "rating": rt, "text": txt,
                "photo_url": "", "published": True, "sort_order": i,
                "created_at": now_iso(), "updated_at": now_iso()})

        faqs = [
            ("Courses", "Do I need prior experience to join?", "No. Most of our courses start from the fundamentals and are designed for both freshers and working professionals."),
            ("Fees", "Do you offer EMI or instalment options?", "Yes, we offer flexible instalment plans and early-bird discounts on most courses."),
            ("Batches", "Do you have weekend batches?", "Yes, we run weekday, weekend and fast-track batches in both online and classroom modes."),
            ("Certification", "Is the certificate recognised?", "You receive a CloudWave Technologies course completion certificate, and we prepare you for global vendor certifications."),
            ("Placement", "Do you guarantee placement?", "We provide dedicated placement assistance including resume building, mock interviews and interview referrals."),
            ("Training", "Are classes live or recorded?", "All classes are live and instructor-led. Recordings are provided for revision."),
        ]
        for i, (cat, q, a) in enumerate(faqs):
            await db.faqs.insert_one({"id": new_id(), "category": cat, "question": q, "answer": a,
                "published": True, "sort_order": i, "created_at": now_iso(), "updated_at": now_iso()})

        blogs = [
            ("Top 5 IT Skills to Learn in 2026", "Career", "The technology landscape keeps evolving. Here are the most in-demand skills that will boost your career this year — cloud computing, automation testing, data analytics, DevOps and full stack development."),
            ("How to Crack Your First Software Testing Interview", "Testing", "Landing your first QA role can feel overwhelming. In this guide we cover the key concepts, common questions and practical tips to help you succeed."),
            ("Why AWS Certification is Worth It in 2026", "Cloud", "Cloud skills are among the highest paid in IT. We break down why an AWS certification can accelerate your career and how to prepare effectively."),
        ]
        for i, (t, cat, content) in enumerate(blogs):
            await db.blog_posts.insert_one({
                "id": new_id(), "slug": slugify(t), "title": t, "category": cat,
                "author": "CloudWave Team", "excerpt": content[:140] + "...",
                "content": content + "\n\n" + content, "image_url": "",
                "tags": [cat, "IT", "Training"], "seo_title": t,
                "seo_description": content[:150], "published": True, "sort_order": i,
                "created_at": now_iso(), "updated_at": now_iso()})

        placements = [
            ("Rahul Sharma", "TCS", "Software Test Engineer", "Selenium with Java", "2026"),
            ("Anjali Gupta", "Infosys", "Cloud Engineer", "AWS Solutions Architect", "2026"),
            ("Karan Mehta", "Wipro", "Data Analyst", "Data Analytics", "2025"),
            ("Pooja Iyer", "Accenture", "Full Stack Developer", "Python Full Stack", "2026"),
            ("Sagar Rao", "Cognizant", "Java Developer", "Java Full Stack", "2025"),
            ("Divya Menon", "Capgemini", "Business Analyst", "Business Analyst", "2026"),
        ]
        for i, (nm, comp, role, crs, yr) in enumerate(placements):
            await db.placements.insert_one({
                "id": new_id(), "student_name": nm, "company_name": comp, "job_title": role,
                "course": crs, "placement_year": yr, "placement_date": f"{yr}-0{(i%9)+1}-15",
                "story": f"{nm} completed the {crs} programme and successfully landed a {role} role at {comp}.",
                "student_photo": "", "company_logo": "", "linkedin_url": "https://linkedin.com",
                "featured": i < 3, "published": True, "sort_order": i,
                "created_at": now_iso(), "updated_at": now_iso()})

        events = [
            ("AWS Cloud Workshop 2026", "Workshop", "A hands-on workshop covering real-world AWS deployment scenarios with industry experts."),
            ("Placement Drive - Selenium Batch", "Placement Drive", "Exclusive placement drive with 15+ hiring partners for our automation testing graduates."),
            ("Career Guidance Seminar", "Seminar", "Free seminar on building a successful IT career, resume tips and interview strategies."),
        ]
        for i, (t, cat, desc) in enumerate(events):
            await db.events.insert_one({
                "id": new_id(), "title": t, "category": cat, "description": desc,
                "event_date": (base + _dt.timedelta(days=15 + i * 12)).isoformat(),
                "location": "CloudWave Center, Pune", "cover_image": "",
                "gallery_images": [], "youtube_url": "", "instagram_url": "",
                "featured": i == 0, "published": True, "sort_order": i,
                "created_at": now_iso(), "updated_at": now_iso()})

        videos = [
            ("Student Success Story - Selenium Batch", "Placements", "https://www.youtube.com/watch?v=dQw4w9WgXcQ"),
            ("AWS Training Demo Session", "Workshops", "https://www.youtube.com/watch?v=dQw4w9WgXcQ"),
            ("Institute Tour & Infrastructure", "Institute Events", "https://www.youtube.com/watch?v=dQw4w9WgXcQ"),
        ]
        for i, (t, cat, url) in enumerate(videos):
            await db.videos.insert_one({
                "id": new_id(), "title": t, "category": cat, "video_url": url,
                "platform": "youtube", "thumbnail_url": "", "description": t,
                "featured": i < 2, "published": True, "sort_order": i,
                "created_at": now_iso(), "updated_at": now_iso()})

        gallery_cats = ["Student Placements", "Certificate Distribution", "Workshops",
                        "Seminars", "Annual Events", "Classroom Activities"]
        for i in range(6):
            await db.gallery.insert_one({
                "id": new_id(), "title": f"{gallery_cats[i]} 2026",
                "description": f"Highlights from our {gallery_cats[i].lower()}.",
                "image_url": "", "category": gallery_cats[i], "event_name": gallery_cats[i],
                "year": "2026", "event_date": "2026-01-15", "location": "Pune",
                "featured": i < 3, "published": True, "sort_order": i,
                "created_at": now_iso(), "updated_at": now_iso()})

    if await db.certificates.count_documents({}) == 0:
        certs = [
            ("CWT-2026-0001", "Rahul Sharma", "Selenium with Java", "2026-03-15", "A+"),
            ("CWT-2026-0002", "Anjali Gupta", "AWS Solutions Architect", "2026-04-10", "A"),
            ("CWT-2026-0003", "Karan Mehta", "Data Analytics", "2026-02-28", "A+"),
        ]
        for i, (cid, nm, crs, dt, grade) in enumerate(certs):
            await db.certificates.insert_one({
                "id": new_id(), "certificate_id": cid, "student_name": nm, "course": crs,
                "issue_date": dt, "grade": grade, "published": True, "sort_order": i,
                "created_at": now_iso(), "updated_at": now_iso()})

    if await db.technologies.count_documents({}) == 0:
        _icon = "https://cdn.simpleicons.org"
        _seed = [
            ("Web Development", "html5", "everyday", "#E34F26"),
            ("React", "react", "everyday", "#61DAFB"),
            ("Node.js", "nodedotjs", "everyday", "#5FA04E"),
            ("AWS", "", "everyday", "#FF9900"),
            ("Docker", "docker", "monday", "#2496ED"),
            ("Kubernetes", "kubernetes", "tuesday", "#326CE5"),
            ("Android Apps", "android", "wednesday", "#3DDC84"),
            ("iOS Apps", "apple", "thursday", "#555555"),
            ("Selenium", "selenium", "friday", "#43B02A"),
            ("Playwright", "playwright", "friday", "#2EAD33"),
            ("Windows Apps", "", "saturday", "#0078D4"),
        ]
        await db.technologies.insert_many([{
            "id": new_id(), "name": nm, "slug_icon": sl,
            "icon_url": (f"{_icon}/{sl}/{col.lstrip('#')}" if sl else ""), "svg_icon": "",
            "color": col, "day_theme": day, "description": "", "active": True,
            "sort_order": i, "created_at": now_iso(), "updated_at": now_iso(),
        } for i, (nm, sl, day, col) in enumerate(_seed)])

    if await db.payment_providers.count_documents({}) == 0:
        await db.payment_providers.insert_one({
            "id": "cloudpay", "name": "CloudPay", "enabled": True, "priority": 100,
            "currencies": ["INR", "USD"], "mode": "sandbox",
            "base_url": ENV_CLOUDPAY["base_url"], "api_key": ENV_CLOUDPAY["api_key"],
            "secret": ENV_CLOUDPAY["secret"],
            "created_at": now_iso(), "updated_at": now_iso()})

    if await db.themes.count_documents({}) == 0:
        await db.themes.insert_one({
            "id": new_id(), "name": "Diwali Celebration", "type": "festival",
            "description": "Festive Diwali accent theme", "start_date": "2026-10-15",
            "end_date": "2026-11-05", "priority": 10, "primary_color": "#F59E0B",
            "secondary_color": "#B45309", "accent_color": "#F59E0B", "background": "",
            "banner_image": "", "decorative_image": "",
            "announcement_text": "🎉 Happy Diwali from CloudWave Technologies! Special admission offers this festive season.",
            "cta_text": "View Celebration Photos", "cta_url": "/gallery",
            "animation_enabled": True, "is_active": True, "sort_order": 0,
            "created_at": now_iso(), "updated_at": now_iso()})
        await db.themes.insert_one({
            "id": new_id(), "name": "New Batch Admission Campaign", "type": "admission_campaign",
            "description": "Admissions open promo", "start_date": "", "end_date": "",
            "priority": 5, "primary_color": "#1D4ED8", "secondary_color": "#4F46E5",
            "accent_color": "#F97316", "background": "", "banner_image": "", "decorative_image": "",
            "announcement_text": "🚀 New batches starting soon! Limited seats — Enquire now for early-bird discounts.",
            "cta_text": "Explore Courses", "cta_url": "/courses",
            "animation_enabled": False, "is_active": True, "sort_order": 1,
            "created_at": now_iso(), "updated_at": now_iso()})


@app.on_event("shutdown")
async def shutdown():
    client.close()
